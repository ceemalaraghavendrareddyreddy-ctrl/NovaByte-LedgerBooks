"""Generic tabular report export — Excel (openpyxl) and PDF (reportlab) — shared by
any report that just needs "the rows on screen, as a file" (aging, and anything
similar added later). Deliberately simple: headers + rows + a title, no per-report
styling. Document-shaped exports (invoices, bills) stay in app/pdf.py.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ACCENT = colors.HexColor("#0F766E")
MUTED = colors.HexColor("#647883")
BORDER = colors.HexColor("#DCE4E9")


def rows_to_xlsx(headers, rows, sheet_title="Report"):
    """rows: list of tuples/lists, same order as headers. Returns a BytesIO ready to send."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Report"  # Excel's own sheet-name length limit

    ws.append(list(headers))
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        cell.fill = header_fill

    for row in rows:
        ws.append(list(row))

    for col_idx, header in enumerate(headers, start=1):
        max_len = max([len(str(header))] + [len(str(row[col_idx - 1])) for row in rows]) if rows else len(str(header))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def rows_to_pdf(title, subtitle, headers, rows, company=None, numeric_cols=None):
    """headers/rows as above; numeric_cols is a set of 0-based column indices to
    right-align (amounts, day counts). Returns a BytesIO ready to send."""
    numeric_cols = numeric_cols or set()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=15 * mm, bottomMargin=15 * mm, leftMargin=15 * mm, rightMargin=15 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Heading1"], textColor=ACCENT, fontSize=16)
    muted_style = ParagraphStyle("Muted", parent=styles["Normal"], textColor=MUTED, fontSize=9)

    elements = [Paragraph(title, title_style)]
    if company and getattr(company, "business_name", None):
        elements.append(Paragraph(company.business_name, muted_style))
    if subtitle:
        elements.append(Paragraph(subtitle, muted_style))
    elements.append(Spacer(1, 8))

    table_data = [list(headers)] + [list(row) for row in rows]
    table = Table(table_data, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.4, BORDER),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFA")]),
    ]
    for col in numeric_cols:
        style.append(("ALIGN", (col, 0), (col, -1), "RIGHT"))
    table.setStyle(TableStyle(style))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer
